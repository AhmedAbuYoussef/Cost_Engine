"""
Excel ↔ engine bridge for the New Corp Model workbook.

* ``extract_state(path)``     – read every *input* cell of the workbook into the
                                state dict consumed by ``cost_engine.compute_all``.
* ``extract_expected(path)``  – read Excel's cached result of every formula cell in
                                the core sheets (the test oracle).
* ``formula_signature(path)`` – the formula text of every core-sheet formula cell;
                                compared against the reference signature to prove a
                                workbook version runs "the same model".

Input cells that hold small constant expressions are decoded explicitly:
``=6176.5/E5`` (an LE amount converted at FX) becomes ``{"le": 6176.5}``,
``=0.07*E6`` (a USD amount converted to LE) becomes ``{"usd": 0.07}``, and pure
arithmetic such as ``=31.718*0.7`` is evaluated.  Anything else in an input cell
raises ``ExtractError`` so a structural change can never be read silently.

Run as a script to (re)generate the JSON fixtures:

    python excel_io.py <workbook.xlsx> <out_dir> [name]
"""

from __future__ import annotations

import ast
import json
import operator
import os
import re
import sys

import openpyxl

COMPANIES = ("EZDK", "EFS", "ERM", "ESR")

CORE_SHEETS = (
    "DRI Cost", "DRI", "Billet", "EZDK Rebar", "EFS Rebar", "ESR Rebar",
    "EZDK Flat", "EFS Flat", "ERM Rolling", "EZDK Wire", "Rebar", "Wire", "Flat",
    "Market Share", "Production Requirments", "Fixed Cost",
    "P&L $ Monthly", "P&L LE Monthly", "P&L $ Annual", "P&L LE Annual",
)


class ExtractError(ValueError):
    pass


# --------------------------------------------------------------------------- #
# Cell decoding
# --------------------------------------------------------------------------- #

_OPS = {ast.Add: operator.add, ast.Sub: operator.sub,
        ast.Mult: operator.mul, ast.Div: operator.truediv}


def _eval_arith(expr: str) -> float:
    """Evaluate +-*/ arithmetic on numeric literals only."""
    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](ev(node.left), ev(node.right))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            v = ev(node.operand)
            return v if isinstance(node.op, ast.UAdd) else -v
        raise ValueError(expr)
    return ev(ast.parse(expr, mode="eval"))


_NUM = r"[0-9]*\.?[0-9]+(?:[eE][-+]?[0-9]+)?"


class _Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.name = ws.title

    def raw(self, cell):
        return self.ws[cell].value

    def where(self, cell):
        return f"'{self.name}'!{cell}"

    def num(self, cell, blank=0.0):
        """Plain numeric input (blank → 0, like Excel); constant arithmetic allowed."""
        v = self.raw(cell)
        if v is None:
            return blank
        if isinstance(v, bool):
            raise ExtractError(f"{self.where(cell)}: unexpected boolean")
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str) and v.startswith("="):
            body = v[1:].replace("$", "")
            try:
                return _eval_arith(body)
            except (ValueError, SyntaxError, ZeroDivisionError):
                pass
        if isinstance(v, str) and v.strip().upper() in ("N/A", "NA", ""):
            return None
        raise ExtractError(f"{self.where(cell)}: expected a number, found {v!r}")

    def money(self, cell, fx_cell):
        """Numeric input that may be written as `=N/<fx>` (LE) or `=N*<fx>` (USD)."""
        v = self.raw(cell)
        if isinstance(v, str) and v.startswith("="):
            body = v[1:].replace("$", "").lstrip("+")
            m = re.fullmatch(rf"\(?({_NUM})\)?/{fx_cell}", body)
            if m:
                return {"le": float(m.group(1))}
            m = re.fullmatch(rf"\(?({_NUM})\)?\*{fx_cell}|{fx_cell}\*\(?({_NUM})\)?", body)
            if m:
                return {"usd": float(m.group(1) or m.group(2))}
        return self.num(cell)

    def label(self, cell):
        v = self.raw(cell)
        return str(v).strip().lower() if v is not None else ""

    def expect_label(self, cell, *fragments):
        lab = self.label(cell)
        if not any(f.lower() in lab for f in fragments):
            raise ExtractError(
                f"{self.where(cell)}: expected a label containing {fragments}, found {lab!r}"
                " — the workbook layout differs from the reference model")


# --------------------------------------------------------------------------- #
# Per-template readers
# --------------------------------------------------------------------------- #

def _read_integrated(s: _Sheet) -> dict:
    """EZDK Rebar / EFS Rebar / EZDK Flat / EFS Flat — EAF → caster → mill."""
    s.expect_label("C30", "byproduct")
    s.expect_label("C109", "byproduct")
    s.expect_label("C163", "byproduct")
    s.expect_label("C186", "work roll")
    n = s.num
    return {
        "stage1": {
            "blend_pct": {"home_scrap": n("D27"), "pig_iron": n("D28")},
            "byproduct_pct_of_solid_charge": n("D30"),
            "consumption": {
                "aux_materials_kg": n("E32"), "refractories_kg": n("E33"),
                "electrodes_kg": n("E34"), "other_fillers_kg": n("E35"),
                "electricity_aux_kwh": n("E38"), "natural_gas_nm3": n("E39"),
                "water_m3": n("E40"), "oxygen_nm3": n("E41"), "nitrogen_nm3": n("E42"),
                "argon_nm3": n("E43"), "handling_kg": n("E45"), "cutting_kg": n("E46"),
            },
            "prices_usd": {
                "home_scrap_t": n("E52"), "pig_iron_t": n("E53"), "byproduct_t": n("E55"),
                "aux_materials_kg": n("E57"), "refractories_kg": n("E58"),
                "electrodes_kg": n("E59"), "other_fillers_kg": n("E60"),
                "natural_gas_nm3": n("E63"), "water_m3": n("E64"), "oxygen_nm3": n("E65"),
                "nitrogen_nm3": n("E66"), "argon_nm3": n("E67"),
                "handling_t": n("E69"), "cutting_t": n("E70"),
            },
        },
        "stage2": {
            "byproduct_pct_of_molten_steel": n("D109"),
            "consumption": {
                "aux_materials_kg": n("E111"), "refractories_kg": n("E112"),
                "electricity_kwh": n("E114"), "natural_gas_nm3": n("E115"),
                "water_m3": n("E116"), "oxygen_nm3": n("E117"),
                "nitrogen_nm3": n("E118"), "argon_nm3": n("E119"),
            },
            "prices_usd": {
                "byproduct_t": n("E124"), "aux_materials_kg": n("E126"),
                "refractories_kg": n("E127"), "natural_gas_nm3": n("E130"),
                "water_m3": n("E131"), "oxygen_nm3": n("E132"),
                "nitrogen_nm3": n("E133"), "argon_nm3": n("E134"),
            },
        },
        "stage3": {
            "feed_share": {"produced": n("D158"), "local": n("D159"), "imported": n("D160")},
            "byproduct_pct_of_feed": n("D163"),
            "consumption": {
                "refractories_kg": n("E165"), "electricity_kwh": n("E167"),
                "natural_gas_nm3": n("E168"), "water_m3": n("E169"),
                "work_roll_units": n("E171"),
            },
            "prices_usd": {
                "local_feed_t": n("E175"), "imported_feed_t": n("E176"),
                "byproduct_t": n("E178"), "refractories_kg": n("E180"),
                "work_roll_unit": n("E186"),
            },
        },
    }


def _read_esr(s: _Sheet) -> dict:
    """ESR Rebar — EAF+BCCM folded into one stage, then rolling."""
    s.expect_label("C32", "byproduct")
    s.expect_label("C111", "byproduct")
    s.expect_label("C140", "tying")
    n, m = s.num, (lambda c: s.money(c, "E5"))
    elec = s.raw("E39")
    mt = re.fullmatch(rf"=\+?Billet!\$?G\$?20(?:\+({_NUM}))?", str(elec).replace(" ", ""))
    if not mt:
        raise ExtractError(f"{s.where('E39')}: expected '=Billet!G20+<extra kWh>', found {elec!r}")
    return {
        "stage1": {
            "blend_pct": {"home_scrap": n("D29"), "pig_iron": n("D30")},
            "byproduct_pct_of_solid_charge": n("D32"),
            "electricity_extra_kwh": float(mt.group(1) or 0.0),
            "consumption": {
                "aux_materials_kg": n("E34"), "refractories_kg": n("E35"),
                "electrodes_kg": n("E36"), "consumables_kg": n("E37"),
                "natural_gas_nm3": n("E40"), "water_m3": n("E41"), "oxygen_nm3": n("E42"),
                "nitrogen_nm3": n("E43"), "chemicals_units": n("E44"),
                "handling_units": n("E46"), "cutting_units": n("E47"),
            },
            "prices_usd": {
                "home_scrap_t": n("E53"), "pig_iron_t": n("E54"), "byproduct_t": m("E56"),
                "aux_materials_kg": m("E58"), "refractories_kg": m("E59"),
                "electrodes_kg": m("E60"), "consumables_kg": m("E61"),
                "natural_gas_nm3": m("E64"), "water_m3": m("E65"), "oxygen_nm3": m("E66"),
                "nitrogen_nm3": m("E67"), "chemicals_unit": m("E68"),
                "handling_unit": m("E70"), "cutting_unit": m("E71"),
            },
        },
        "stage3": {
            "feed_share": {"produced": n("D106"), "local": n("D107"), "imported": n("D108")},
            "byproduct_pct_of_feed": n("D111"),
            "crops_t": n("E112"),
            "consumption": {
                "refractories_kg": n("E114"), "electricity_kwh": n("E116"),
                "natural_gas_nm3": n("E117"), "water_m3": n("E118"),
                "labels_rings_units": n("E120"), "chemicals_units": n("E121"),
                "tying_kg": n("E122"),
            },
            "prices_usd": {
                "imported_feed_t": n("E127"), "byproduct_t": n("E129"), "crops_t": n("E130"),
                "refractories_kg": n("E132"), "labels_rings_unit": m("E138"),
                "chemicals_unit": m("E139"), "tying_t": m("E140"),
            },
        },
    }


def _read_rolling(s: _Sheet) -> dict:
    """ERM Rolling / EZDK Wire — rolling only, built in LE."""
    s.expect_label("C27", "by-product", "byproduct")
    s.expect_label("C44", "by product")
    n, m = s.num, (lambda c: s.money(c, "E6"))
    bp = s.raw("E44")
    byproduct = "billet_price" if isinstance(bp, str) and re.fullmatch(r"=\+?\$?E\$?42", bp) else n("E44")
    return {
        "feed_share": {"imported": n("D24"), "produced": n("D25"), "local": n("D26")},
        "byproduct_pct_of_feed": n("D27"),
        "consumption": {
            "electricity_kwh": n("E28"), "natural_gas_nm3": n("E29"),
            "recycled_water_nm3": n("E31"), "tying_kg": n("E32"), "chemicals_kg": n("E33"),
            "labels_rings_kg": n("E34"), "spare_parts_kg": n("E35"),
            "external_services_kg": n("E36"),
        },
        "prices_le": {
            "imported_billet_usd_t": n("E39"), "imported_extra_le_t": n("E40"),
            "local_billet_t": n("E43"), "byproduct_t": byproduct,
            "electricity_kwh": m("E45"), "natural_gas_nm3": m("E47"),
            "recycled_water_nm3": m("E48"), "tying_kg": m("E49"), "chemicals_kg": m("E50"),
            "labels_rings_kg": m("E51"), "spare_parts_kg": m("E52"),
            "external_services_kg": m("E53"),
        },
    }


def _read_dri(dri: _Sheet, cost: _Sheet, cost_values) -> dict:
    out = {}
    for co, c_dri, c_cost in (("EZDK", "E", "E"), ("ERM", "F", "G")):
        n = cost.num
        other = cost.raw(f"{c_cost}49")
        mt = re.fullmatch(rf"=({_NUM})-({_NUM})-\$?{c_cost}\$?47", str(other).replace(" ", ""))
        if mt:
            other_fixed = {"total_le": float(mt.group(1)), "less_depreciation_le": float(mt.group(2))}
        else:
            other_fixed = n(f"{c_cost}49")
        out[co] = {
            "iop_landed_usd_t": dri.num(f"{c_dri}8"),
            "ng_price_usd_mmbtu": dri.num(f"{c_dri}9"),
            "mrmr": dri.num(f"{c_dri}11"),
            "capacity_t": n(f"{c_cost}10"),
            "capacity_utilization": n(f"{c_cost}11"),
            # Gross margin (row 87) uses row 86, a link to an external workbook
            # ('[1]Production Q. & Sales (2)'); row 13 is never referenced.
            **_dri_selling_price(cost, cost_values, c_cost),
            "consumption": {
                "electricity_kwh": n(f"{c_cost}20"), "natural_gas_nm3": dri.num(f"{c_dri}12"),
                "oxygen_nm3": n(f"{c_cost}23"), "nitrogen_nm3": n(f"{c_cost}24"),
                "water_nm3": n(f"{c_cost}25"), "chemicals_t": n(f"{c_cost}26"),
                "spare_parts_t": n(f"{c_cost}27"), "external_services_t": n(f"{c_cost}28"),
                "other_t": n(f"{c_cost}29"),
            },
            "prices_le": {
                "electricity_kwh": n(f"{c_cost}35"), "oxygen_nm3": n(f"{c_cost}38"),
                "nitrogen_nm3": n(f"{c_cost}39"), "water_nm3": n(f"{c_cost}40"),
                "chemicals_t": n(f"{c_cost}41"), "spare_parts_t": n(f"{c_cost}42"),
                "external_services_t": n(f"{c_cost}43"), "other_t": n(f"{c_cost}44"),
            },
            "fixed_le": {
                "labor": n(f"{c_cost}47"), "depreciation": n(f"{c_cost}48"),
                "other_fixed": other_fixed,
            },
        }
    return out


def _dri_selling_price(cost: _Sheet, cost_values, col: str) -> dict:
    cached = cost_values[f"{col}86"].value
    if isinstance(cached, (int, float)) and not isinstance(cached, bool):
        return {"selling_price_le_t": float(cached),
                "selling_price_source": f"cached value of external link 'DRI Cost'!{col}86",
                "selling_price_input_le_t": cost.num(f"{col}13")}
    return {"selling_price_le_t": cost.num(f"{col}13"),
            "selling_price_source": f"'DRI Cost'!{col}13 (external link {col}86 unavailable)"}


def _read_billet(s: _Sheet) -> dict:
    out = {}
    for co, c in (("EZDK", "E"), ("EFS", "F"), ("ESR", "G")):
        scrap = {}
        for key, row in (("local_scrap_usd_t", 9), ("imported_scrap_usd_t", 10)):
            v = s.raw(f"{c}{row}")
            mt = re.fullmatch(rf"=\+?\$?([EFG])\$?{row}", str(v)) if isinstance(v, str) else None
            scrap[key] = {"same_as": {"E": "EZDK", "F": "EFS", "G": "ESR"}[mt.group(1)]} if mt else s.num(f"{c}{row}")
        out[co] = {
            "dri_margin_from_erm_usd_t": s.num(f"{c}6") if co != "EZDK" else 0.0,
            **scrap,
            "electricity_usd_kwh": s.num(f"{c}11"),
            "blend_pct": {"dri": s.num(f"{c}13"), "local_scrap": s.num(f"{c}14"),
                          "imported_scrap": s.num(f"{c}15")},
            "eaf_yield_pct": s.num(f"{c}17"),
            "ccp_yield_pct": s.num(f"{c}18"),
            "eaf_lf_electricity_kwh_t_ms": s.num(f"{c}20"),
        }
    out["EZDK"]["tradeoff_ratio"] = s.num("E35")
    out["EFS"]["tradeoff_ratio"] = s.num("F35")
    out["ESR"]["tradeoff_ratio"] = s.num("H35")
    overrides: dict = {}
    for r, src in ((39, "EZDK"), (40, "EFS"), (41, "ESR")):
        for c, buyer in (("E", "EZDK"), ("F", "EFS"), ("G", "ERM"), ("H", "ESR")):
            v = s.raw(f"{c}{r}")
            if not (isinstance(v, str) and v.startswith("=")):
                overrides.setdefault(src, {})[buyer] = s.num(f"{c}{r}")
    return {
        "companies": out,
        "landed_cost_factor": s.num("E36"),
        "market_usd_t": {"base": s.num("E48"), "safe_guards": s.num("E49"),
                         "other_costs": s.num("E50")},
        "tradeoff_overrides": overrides,
    }


def _read_flat(s: _Sheet) -> dict:
    out = {}
    for co, c in (("EZDK", "E"), ("EFS", "F")):
        out[co] = {
            "dri_margin_from_erm_usd_t": s.num(f"{c}6") if co != "EZDK" else 0.0,
            "local_scrap_usd_t": s.num(f"{c}9"),
            "imported_scrap_usd_t": s.num(f"{c}10"),
            "electricity_usd_kwh": s.num(f"{c}11"),
            "blend_pct": {"dri": s.num(f"{c}13"), "local_scrap": s.num(f"{c}14"),
                          "imported_scrap": s.num(f"{c}15")},
            "eaf_yield_pct": s.num(f"{c}17"),
            "tsc_yield_pct": s.num(f"{c}18"),
            "hsm_yield_pct": s.num(f"{c}19"),
            "eaf_lf_electricity_kwh_t_ms": s.num(f"{c}20"),
        }
    return out


def _read_sales(ms: _Sheet, pl: _Sheet) -> tuple[dict, dict, dict]:
    ms.expect_label("B19", "ezdk export")
    ms.expect_label("B20", "efs export")
    n = ms.num
    sales = {
        "EZDK": {
            "Rebar": {"local_kt": n("E6"), "export_kt": n("E19"),
                      "local_price_le_t": n("E29"), "export_price_usd_t": n("E34")},
            "Wire Rod": {"local_kt": n("F6"), "export_kt": n("F19"),
                         "local_price_le_t": n("F29"), "export_price_usd_t": n("F34")},
            "HRC": {"local_kt": n("G6"), "export_kt": n("G19"),
                    "local_price_le_t": n("G29"), "export_price_usd_t": n("G34")},
        },
        "EFS": {
            "Rebar": {"local_kt": n("E7"), "export_kt": n("E20"),
                      "local_price_le_t": n("E30"), "export_price_usd_t": n("E35")},
            "HRC": {"local_kt": n("G7"), "export_kt": n("G20"),
                    "local_price_le_t": n("G30"), "export_price_usd_t": n("G35")},
        },
        "ERM": {"Rebar": {"local_kt": n("E8"), "export_kt": 0.0,
                          "local_price_le_t": n("E31"), "export_price_usd_t": n("E36")}},
        "ESR": {"Rebar": {"local_kt": n("E9"), "export_kt": 0.0,
                          "local_price_le_t": n("E32"), "export_price_usd_t": n("E37")}},
    }
    market = {"Rebar": n("E22", blank=None), "Wire Rod": n("F22", blank=None),
              "HRC": n("G22", blank=None)}
    p = pl.num
    pnl = {
        "export_expense_rate": {
            "EZDK": {"Rebar": p("E21"), "Wire Rod": p("F21"), "HRC": p("G21")},
            "EFS": {"Rebar": p("J21"), "HRC": p("K21")},
            "ERM": {"Rebar": p("N21")},
            "ESR": {"Rebar": p("P21")},
        },
        "intercompany_gain_dri_musd": {"ERM": p("N28")},
    }
    return sales, market, pnl


def _read_fixed(s: _Sheet) -> tuple[float, dict]:
    s.expect_label("B15", "exchange")
    out = {}
    for co, c in zip(COMPANIES, "DEFG"):
        out[co] = {
            "manufacturing_musd": s.num(f"{c}7"), "sga_musd": s.num(f"{c}8"),
            "net_finance_musd": s.num(f"{c}9"), "depreciation_musd": s.num(f"{c}11"),
            "distribution_pct": {"DRI": s.num(f"{c}30"), "Rebar": s.num(f"{c}31"),
                                 "Wire Rod": s.num(f"{c}32"), "HRC": s.num(f"{c}33")},
        }
    return s.num("D15"), out


def _read_capacity(S: dict) -> dict:
    """Monthly ceilings (t/month) from the workbook's annual capacity cells."""
    for sheet, cell in (("DRI Cost", "C10"), ("EZDK Rebar", "C8"), ("EZDK Wire", "C11"),
                        ("EZDK Flat", "C8"), ("ERM Rolling", "C11")):
        S[sheet].expect_label(cell, "capacity", "production cap")
    m = lambda sheet, cell: S[sheet].num(cell) / 12
    return {
        "dri": {"EZDK": m("DRI Cost", "E10"), "ERM": m("DRI Cost", "G10")},
        "finished": {
            "EZDK": {"Rebar": m("EZDK Rebar", "E8"), "Wire Rod": m("EZDK Wire", "E11"),
                     "HRC": m("EZDK Flat", "E8")},
            "EFS": {"Rebar": m("EFS Rebar", "E8"), "HRC": m("EFS Flat", "E8")},
            "ERM": {"Rebar": m("ERM Rolling", "E11")},
            "ESR": {"Rebar": m("ESR Rebar", "E8")},
        },
    }


def _detect_erm_billet_source(rebar: _Sheet) -> str:
    v = str(rebar.raw("G12")).replace("$", "")
    mt = re.fullmatch(r"=\+?Billet!([EFG])30", v)
    if mt:
        return "vc:" + {"E": "EZDK", "F": "EFS", "G": "ESR"}[mt.group(1)]
    if re.fullmatch(r"=\+?Billet!E52", v):
        return "market"
    raise ExtractError(f"'Rebar'!G12: unrecognised ERM billet source {v!r}")


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #

def extract_state(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    missing = [s for s in CORE_SHEETS if s not in wb.sheetnames]
    if missing:
        raise ExtractError(f"workbook is missing core sheets: {missing}")
    S = {name: _Sheet(wb[name]) for name in CORE_SHEETS}
    fx, fixed = _read_fixed(S["Fixed Cost"])
    sales, market, pnl = _read_sales(S["Market Share"], S["P&L $ Monthly"])
    return {
        "meta": {"source_workbook": os.path.basename(path)},
        # an extracted workbook is evaluated with the workbook's own (legacy) rules;
        # model_fixes.to_corrected() turns it into a corrected-model state
        "rules_profile": "legacy",
        "fx_egp_per_usd": fx,
        "sales": sales,
        "total_local_market_kt": market,
        "dri": _read_dri(S["DRI"], S["DRI Cost"], wv["DRI Cost"]),
        "billet": _read_billet(S["Billet"]),
        "flat": _read_flat(S["Flat"]),
        "finishing_yield_pct": {
            "Rebar": {co: S["Rebar"].num(f"{c}6") for co, c in zip(COMPANIES, "EFGH")},
            "Wire Rod": {"EZDK": S["Wire"].num("E6")},
        },
        "sourcing": {"ERM_rebar_billet": _detect_erm_billet_source(S["Rebar"])},
        "detail": {
            "EZDK Rebar": _read_integrated(S["EZDK Rebar"]),
            "EFS Rebar": _read_integrated(S["EFS Rebar"]),
            "EZDK Flat": _read_integrated(S["EZDK Flat"]),
            "EFS Flat": _read_integrated(S["EFS Flat"]),
            "ESR Rebar": _read_esr(S["ESR Rebar"]),
            "ERM Rolling": _read_rolling(S["ERM Rolling"]),
            "EZDK Wire": _read_rolling(S["EZDK Wire"]),
        },
        "fixed_cost": fixed,
        "pnl": pnl,
        "capacity_ceilings": _read_capacity(S),
    }


def extract_expected(path: str) -> dict:
    """Cached numeric value of every formula cell in the core sheets: {'Sheet!A1': v}."""
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in CORE_SHEETS:
        for row in wf[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = wv[name][c.coordinate].value
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        out[f"{name}!{c.coordinate}"] = float(v)
    return out


def formula_signature(path: str) -> dict:
    wf = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for name in CORE_SHEETS:
        for row in wf[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[f"{name}!{c.coordinate}"] = c.value
    return out


_IGNORED_FORMULA_CELLS = re.compile(
    r"!(H2|G2|F2|E2|I2|K2)$"            # date stamps / title links
    r"|^(Billet!H54|Rebar!H36|Wire!E36|Flat!F32|Market Share!E39|DRI!G24|"
    r"Production Requirments!H47|P&L .*!H50)$"   # 'n / 12' page counters
)


def _is_input_formula(f) -> bool:
    """A constant typed as a formula (e.g. '=31.7*0.7', '=6176.5/E5', '=0.07*E6',
    '=347999806.05-239746417.14-E47'): its numbers are input figures, not model logic."""
    if not isinstance(f, str):
        return False
    body = re.sub(r"\$?E\$?(5|6|47)\b", "", f[1:].replace("$", ""))
    return re.fullmatch(r"[-+*/(). 0-9eE]*", body) is not None and any(ch.isdigit() for ch in body)


def check_same_model(path: str, reference_path: str) -> list[str]:
    """Core-sheet formula cells whose text differs from the reference workbook.

    An empty list means ``path`` runs the same formulas as the reference model,
    so the engine applies to it unchanged; anything listed needs a look before
    the engine's numbers can be trusted for that file."""
    a, b = formula_signature(path), formula_signature(reference_path)
    diffs = []
    for cell in sorted(set(a) | set(b)):
        if _IGNORED_FORMULA_CELLS.search(cell):
            continue
        fa, fb = a.get(cell), b.get(cell)
        if fa == fb or (_is_input_formula(fa) and _is_input_formula(fb)):
            continue
        diffs.append(f"{cell}: {b.get(cell)!r} → {a.get(cell)!r}")
    return diffs


def main(argv):
    if len(argv) < 3:
        print(__doc__)
        return 2
    path, out_dir = argv[1], argv[2]
    name = argv[3] if len(argv) > 3 else os.path.splitext(os.path.basename(path))[0]
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, f"{name}_state.json"), "w") as f:
        json.dump(extract_state(path), f, indent=1)
    with open(os.path.join(out_dir, f"{name}_expected.json"), "w") as f:
        json.dump(extract_expected(path), f, indent=0, sort_keys=True)
    print(f"wrote {name}_state.json and {name}_expected.json to {out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
