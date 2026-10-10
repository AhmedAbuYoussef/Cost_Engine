"""
Monthly baseline refresh: the seven-sheet XLSX template (system prompt, Appendix D.5)
and the parsers that turn a filled template — or a CSV — into a complete state.

Every input of the state appears on exactly one sheet, one row per input, with three
value columns: "Last month (read-only)", "This month" (blank by default) and
"% change".  A blank "This month" cell means unchanged: the parser carries last month's
value forward, so the result is always a full snapshot (Tool 13 commit discipline).

Column A of each data sheet holds the input's path; it is how the parser finds the input
again, so rows can be re-ordered or filtered freely.
"""

from __future__ import annotations

import copy
import csv
import json

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

import state_schema as schema
from state_schema import ModelError

SHEETS = ("Cover & Metadata", "DRI Inputs", "Billet Inputs", "Finished Products", "Sales",
          "Fixed Costs & Capacity", "Intercompany")
HEADER = ("Path", "Company / line", "Input", "Unit", "Last month (read-only)", "This month",
          "% change")
COL_PATH, COL_LAST, COL_THIS = 1, 5, 6

_BILLET_DETAIL = ("detail.EZDK Rebar.stage1", "detail.EZDK Rebar.stage2", "detail.EFS Rebar.stage1",
                  "detail.EFS Rebar.stage2", "detail.ESR Rebar.stage1")


def sheet_for(path: str) -> str | None:
    """Which template sheet an input belongs to (None = managed by the system)."""
    root = path.split(".")[0]
    if root in schema.READ_ONLY_ROOTS or path == "billet.tradeoff_overrides":
        return None
    if root == "fx_egp_per_usd":
        return "Cover & Metadata"
    if root == "dri":
        return "DRI Inputs"
    if root == "sourcing" or path.endswith("dri_margin_from_erm_usd_t"):
        return "Intercompany"
    if root == "billet" or path.startswith(_BILLET_DETAIL):
        return "Billet Inputs"
    if root in ("detail", "flat", "finishing_yield_pct"):
        return "Finished Products"
    if root in ("sales", "pnl", "total_local_market_kt"):
        return "Sales"
    if root in ("fixed_cost", "capacity_ceilings", "policy_floors"):
        return "Fixed Costs & Capacity"
    raise ModelError("schema_validation_error", f"no template sheet for input {path!r}", field=path)


def _describe(path: str) -> tuple:
    """(company / line, input name, unit) for a row."""
    segs = path.split(".")
    who = next((s for s in segs if s in schema.COMPANIES or s.split(" ")[0] in schema.COMPANIES), "Group")
    rest = [s for s in segs[1:] if s != who and s != "companies"]
    name = " · ".join(s.replace("_", " ") for s in rest) or segs[0].replace("_", " ")
    key = segs[-1]
    unit = ("%" if key.endswith("_pct") or ".blend_pct." in path or ".distribution_pct." in path
            else "t/month" if path.startswith("capacity_ceilings") else
            "kt" if key.endswith("_kt") else
            "LE/t" if key.endswith("_le_t") else "$/t" if key.endswith("_usd_t") else "")
    return who, name, unit


def _leaf_value(v):
    """(number or text shown in the template, unit override)."""
    if isinstance(v, dict) and len(v) == 1:
        k, x = next(iter(v.items()))
        return x, ("in LE (converted at FX)" if k == "le" else "in USD (converted at FX)")
    return v, None


def generate(state: dict, path: str, *, source_label: str, pre_fill_this_month: bool = False) -> str:
    """Write the refresh template for ``state`` (last month) to ``path``."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold, grey = Font(bold=True), PatternFill("solid", fgColor="EEEEEE")
    rows = {name: [] for name in SHEETS}
    for p, v in schema.leaves(state).items():
        sheet = sheet_for(p)
        if sheet:
            rows[sheet].append((p, v))

    cover = wb.create_sheet("Cover & Metadata")
    cover.append(["Ezz Steel model — monthly baseline refresh"])
    cover["A1"].font = Font(bold=True, size=14)
    cover.append([f"Last month: {source_label} (month {state.get('month_label')})"])
    cover.append([])
    meta = [("Month label (YYYY-MM)", None), ("Author", None), ("Notes", None)]
    for label, _ in meta:
        cover.append([label, ""])
        cover.cell(cover.max_row, 1).font = bold
    cover.append([])
    cover.append(["Sections changing this month (Yes / No)"])
    cover.cell(cover.max_row, 1).font = bold
    for name in SHEETS[1:]:
        cover.append([name, ""])
    cover.append([])
    cover.column_dimensions["A"].width = 44
    cover.column_dimensions["B"].width = 30

    for name in SHEETS:
        ws = cover if name == "Cover & Metadata" else wb.create_sheet(name)
        start = ws.max_row + 1 if name == "Cover & Metadata" else 1
        for c, h in enumerate(HEADER, 1):
            ws.cell(start, c, h).font = bold
            ws.cell(start, c).fill = grey
        if name == "Billet Inputs":
            ws.cell(start, 9, "ERM has no furnace — it buys billets, so it has no rows here.")
        for p, v in rows[name]:
            who, label, unit = _describe(p)
            shown, unit_note = _leaf_value(v)
            r = ws.max_row + 1
            ws.cell(r, 1, p)
            ws.cell(r, 2, who)
            ws.cell(r, 3, label)
            ws.cell(r, 4, unit_note or unit)
            ws.cell(r, COL_LAST, shown)
            if pre_fill_this_month:
                ws.cell(r, COL_THIS, shown)
            if isinstance(shown, (int, float)) and not isinstance(shown, bool):
                L, T = get_column_letter(COL_LAST), get_column_letter(COL_THIS)
                ws.cell(r, 7, f'=IF(OR({T}{r}="",{L}{r}=0),"",{T}{r}/{L}{r}-1)').number_format = "0.0%"
            if p == "total_local_market_kt.HRC":
                ws.cell(r, 8, "Left blank by design: HRC market size is asked at each run.")
            ws.cell(r, COL_LAST).protection = Protection(locked=True)
        _add_checks(ws, name, rows[name])
        ws.column_dimensions["A"].hidden = True
        for col, width in zip("BCDEFG", (16, 58, 24, 20, 16, 10)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = ws.cell(start + 1, 2)
        for row in ws.iter_rows(min_row=start + 1):
            for c in row:
                c.alignment = Alignment(vertical="top")
    wb.save(path)
    return path


def _add_checks(ws, name, rows):
    """Sum-to-100 check cells (blends per furnace, fixed-cost distribution per company)."""
    groups = {}
    for i, (p, _) in enumerate(rows):
        if name == "Billet Inputs" and p.startswith("billet.companies.") and ".blend_pct." in p:
            groups.setdefault(f"Blend total {p.split('.')[2]} long (must be 100)", []).append(p)
        if name == "Finished Products" and p.startswith("flat.") and ".blend_pct." in p:
            groups.setdefault(f"Blend total {p.split('.')[1]} flat (must be 100)", []).append(p)
        if name == "Fixed Costs & Capacity" and ".distribution_pct." in p:
            groups.setdefault(f"Distribution total {p.split('.')[1]} (must be 100)", []).append(p)
    if not groups:
        return
    index = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    ws.append([])
    L, T = get_column_letter(COL_LAST), get_column_letter(COL_THIS)
    for label, paths in groups.items():
        terms = "+".join(f'IF({T}{index[p]}="",{L}{index[p]},{T}{index[p]})' for p in paths)
        r = ws.max_row + 1
        ws.cell(r, 3, label).font = Font(bold=True)
        ws.cell(r, COL_THIS, f"={terms}")
        ws.cell(r, 7, f'=IF(ABS({T}{r}-100)<0.01,"OK","CHECK")')


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #

def _coerce(base_leaf, raw, path):
    """A template / CSV cell → a value of the same kind as last month's."""
    if isinstance(raw, str):
        raw = raw.strip()
        if raw == "":
            return None
    if isinstance(base_leaf, str):
        return str(raw)
    try:
        x = float(raw)
    except (TypeError, ValueError):
        raise ModelError("schema_validation_error", f"{path}: {raw!r} is not a number", field=path,
                         expected="a number")
    if isinstance(base_leaf, dict) and len(base_leaf) == 1:
        return {next(iter(base_leaf)): x}
    return x


def _apply(base: dict, values: dict) -> tuple:
    """Last month + this month's entries → (full payload, [changes])."""
    payload = copy.deepcopy(base)
    leaves = schema.leaves(base)
    changes = []
    for path, raw in values.items():
        if path not in leaves or sheet_for(path) is None:
            if schema.exists(base, path) or path.split(".")[0] in schema.READ_ONLY_ROOTS:
                raise ModelError("invalid_path", f"{path!r} cannot be set by a refresh", field=path)
            raise schema.missing_path_error(path)
        new = _coerce(leaves[path], raw, path)
        if new is None:
            continue
        if json.dumps(new) != json.dumps(leaves[path]):
            schema.set_value(payload, path, new)
            changes.append({"path": path, "last_month": leaves[path], "this_month": new})
    return payload, changes


def parse_xlsx(path: str, base: dict) -> dict:
    """Read a filled template.  Returns {payload, changes, month_label, author, notes}."""
    wb = openpyxl.load_workbook(path, data_only=False)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise ModelError("schema_validation_error", f"not a refresh template — missing sheets {missing}",
                         field="file_path")
    values, meta = {}, {}
    for name in SHEETS:
        ws = wb[name]
        for r in range(1, ws.max_row + 1):
            p = ws.cell(r, COL_PATH).value
            label = ws.cell(r, 1).value
            if name == "Cover & Metadata" and label in ("Month label (YYYY-MM)", "Author", "Notes"):
                meta[label] = ws.cell(r, 2).value
                continue
            if not p or p == "Path" or not isinstance(p, str) or "." not in p and p != "fx_egp_per_usd":
                continue
            cell = ws.cell(r, COL_THIS).value
            if cell is not None and not (isinstance(cell, str) and cell.startswith("=")):
                values[p] = cell
    payload, changes = _apply(base, values)
    month = meta.get("Month label (YYYY-MM)")
    return {"payload": payload, "changes": changes,
            "month_label": str(month).strip() if month else None,
            "author": meta.get("Author"), "notes": meta.get("Notes")}


def parse_csv(path: str, base: dict) -> dict:
    """CSV with a header row ``path,value`` (extra columns are ignored)."""
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames or not {"path", "value"} <= {h.strip().lower() for h in reader.fieldnames}:
            raise ModelError("schema_validation_error", "the CSV needs a header row with 'path' and 'value'",
                             field="file_path", expected="path,value")
        key = {h.strip().lower(): h for h in reader.fieldnames}
        values = {row[key["path"]].strip(): row[key["value"]] for row in reader if row[key["path"]]}
    payload, changes = _apply(base, values)
    return {"payload": payload, "changes": changes, "month_label": None, "author": None, "notes": None}
