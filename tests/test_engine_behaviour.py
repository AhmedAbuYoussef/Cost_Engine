"""Engine behaviour beyond Excel fidelity: purity, integrity checks, the ERM sourcing
switch, and the extractor's refusal to read a structurally different workbook."""

import copy
import json

import openpyxl
import pytest

import cost_engine
import excel_io
from conftest import REFERENCE_WORKBOOK


def test_pure_and_deterministic(state):
    before = copy.deepcopy(state)
    a = cost_engine.compute_all(state)
    b = cost_engine.compute_all(state)
    assert state == before
    assert json.dumps(a, sort_keys=True, default=str) == json.dumps(b, sort_keys=True, default=str)


def test_state_is_json_round_trippable(state):
    again = json.loads(json.dumps(state))
    assert cost_engine.compute_all(again)["pnl"]["usd_monthly"]["columns"]["R"]["ebt"] == \
        cost_engine.compute_all(state)["pnl"]["usd_monthly"]["columns"]["R"]["ebt"]


# ---- integrity checks ------------------------------------------------------ #

def _check(outputs, name):
    return next(c for c in outputs["integrity"] if c["check"] == name)


def test_reference_blending_ratios_flagged(state):
    """The reference workbook's EZDK/EFS blends sum to 99.977–99.978%, not 100%."""
    c = _check(cost_engine.compute_all(state), "blending_100")
    assert not c["passed"]
    assert "EZDK long" in c["detail"] and "EFS long" in c["detail"]
    assert "ESR long" not in c["detail"]


def test_blending_check_passes_when_ratios_sum_to_100(state):
    for co in ("EZDK", "EFS"):
        bl = state["billet"]["companies"][co]["blend_pct"]
        bl["imported_scrap"] = 100 - bl["dri"] - bl["local_scrap"]
    for co in ("EZDK", "EFS"):
        bl = state["flat"][co]["blend_pct"]
        bl["imported_scrap"] = 100 - bl["dri"] - bl["local_scrap"]
    assert _check(cost_engine.compute_all(state), "blending_100")["passed"]


def test_distribution_check_fires(state):
    state["fixed_cost"]["ERM"]["distribution_pct"]["DRI"] = 60
    out = cost_engine.compute_all(state)
    c = _check(out, "fixed_distribution_100")
    assert not c["passed"] and "ERM" in c["detail"]


def test_strict_mode_raises(state):
    with pytest.raises(cost_engine.IntegrityError):
        cost_engine.compute_all(state, strict=True)


@pytest.mark.parametrize("company, product, col", [
    ("ERM", "Rebar", "N"), ("ESR", "Rebar", "P"), ("EFS", "Rebar", "J"), ("EZDK", "Rebar", "E"),
])
def test_zero_sales_product(state, company, product, col):
    """Zero sales is a legitimate scenario (rulebook check 7).  Excel shows #DIV/0!; the
    engine keeps the per-ton costs, reports zero volumes and a break-even of 0."""
    base = cost_engine.compute_all(state)
    state["sales"][company][product]["local_kt"] = 0.0
    state["sales"][company][product]["export_kt"] = 0.0
    out = cost_engine.compute_all(state)
    c = out["pnl"]["usd_monthly"]["columns"][col]
    assert c["total_qty"] == 0 and c["break_even_qty"] == 0
    assert c["avg_price"] is None                       # Excel: #DIV/0!
    assert c["cost_per_t"] == pytest.approx(base["pnl"]["usd_monthly"]["columns"][col]["cost_per_t"])
    assert c["variable_cogs"] == 0
    assert _check(out, "break_even_zero")["passed"]


def test_zero_rebar_keeps_ezdk_wire_billets(state):
    """EZDK's billet shop also feeds wire rod: with no rebar sales it still runs."""
    state["sales"]["EZDK"]["Rebar"]["local_kt"] = 0.0
    out = cost_engine.compute_all(state)
    d = out["detail"]["EZDK Rebar"]
    wire_billets = out["detail"]["EZDK Wire"]["billets_t"]
    assert d["finished_t"] == 0
    assert d["s2.output_t"] == pytest.approx(wire_billets)
    assert d["s3.feed.produced_t"] == pytest.approx(0, abs=1e-6)


def test_every_view_declares_currency(state):
    out = cost_engine.compute_all(state)
    assert out["pnl"]["le_monthly"]["_currency"] == "LE"
    assert out["pnl"]["usd_annual"]["_currency"] == "USD"
    assert out["billet"]["tradeoff"]["_currency"] == "USD"
    assert _check(out, "currency_tagged")["passed"]


# ---- ERM billet sourcing ---------------------------------------------------- #

@pytest.mark.parametrize("source, expected", [
    ("vc:EZDK", lambda o: o["billet"]["EZDK"]["total_variable_cost"]),
    ("vc:ESR", lambda o: o["billet"]["ESR"]["total_variable_cost"]),
    ("market", lambda o: o["billet"]["market_price"]),
    ("min", lambda o: o["billet"]["tradeoff"]["minimum"]["ERM"]),
    ("offer:EFS", lambda o: o["billet"]["tradeoff"]["rows"]["EFS"]["ERM"]),
])
def test_erm_sourcing_dispatch(state, source, expected):
    state["sourcing"]["ERM_rebar_billet"] = source
    out = cost_engine.compute_all(state)
    price = expected(out)
    assert out["rebar"]["ERM"]["sc1"]["material_price"] == pytest.approx(price)
    fx = state["fx_egp_per_usd"]
    assert out["detail"]["ERM Rolling"]["price.produced_billet"] == pytest.approx(price * fx)


def test_reference_erm_source_is_ezdk_vc(reference_state):
    assert reference_state["sourcing"]["ERM_rebar_billet"] == "vc:EZDK"


# ---- cascade behaviour ------------------------------------------------------ #

def test_sales_drive_production(state):
    base = cost_engine.compute_all(state)["production"]
    state["sales"]["ESR"]["Rebar"]["local_kt"] *= 2
    after = cost_engine.compute_all(state)["production"]
    assert after["long"]["ESR"]["billet"] == pytest.approx(2 * base["long"]["ESR"]["billet"])
    # ESR's DRI comes from ERM, so ERM's IOP requirement rises too
    assert after["long"]["ERM"]["iop"] > base["long"]["ERM"]["iop"]


def test_fx_flows_to_local_prices(state):
    base = cost_engine.compute_all(state)["pnl"]["usd_monthly"]["columns"]["E"]["local_price"]
    state["fx_egp_per_usd"] *= 2
    after = cost_engine.compute_all(state)["pnl"]["usd_monthly"]["columns"]["E"]["local_price"]
    assert after == pytest.approx(base / 2)


# ---- extractor safety --------------------------------------------------------- #

def test_extractor_rejects_formula_in_input_cell(tmp_path):
    wb = openpyxl.load_workbook(REFERENCE_WORKBOOK)
    wb["EZDK Rebar"]["E32"] = "=E31*2"
    path = tmp_path / "broken.xlsx"
    wb.save(path)
    with pytest.raises(excel_io.ExtractError, match="EZDK Rebar"):
        excel_io.extract_state(str(path))


def test_extractor_rejects_moved_rows(tmp_path):
    wb = openpyxl.load_workbook(REFERENCE_WORKBOOK)
    wb["EZDK Rebar"]["C30"] = "Something else"
    path = tmp_path / "moved.xlsx"
    wb.save(path)
    with pytest.raises(excel_io.ExtractError, match="layout differs"):
        excel_io.extract_state(str(path))


def test_same_model_check(tmp_path):
    wb = openpyxl.load_workbook(REFERENCE_WORKBOOK)
    wb["EFS Rebar"]["E168"] = "=31.7184290558378*0.6"     # a different input figure: fine
    wb["Billet"]["E24"] = "=E8*E13%+E9*E14%"              # a different formula: flagged
    path = tmp_path / "variant.xlsx"
    wb.save(path)
    diffs = excel_io.check_same_model(str(path), REFERENCE_WORKBOOK)
    assert len(diffs) == 1 and diffs[0].startswith("Billet!E24")
