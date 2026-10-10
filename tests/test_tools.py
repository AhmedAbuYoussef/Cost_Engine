"""Step 3 — the 14 tools.  Done-criterion (Appendix C): each tool callable directly from
Python with structured arguments and structured returns.  Every number a tool shows must
be the engine's number."""

import csv
import json

import openpyxl
import pytest

import cost_engine as ce
import refresh_template as rt
import state_manager as sm
import state_schema as schema
import tools
from conftest import REFERENCE_WORKBOOK


@pytest.fixture
def store(tmp_path):
    s = sm.StateStore(str(tmp_path / "m.db"))
    yield s
    s.close()


@pytest.fixture
def box(store, tmp_path):
    return tools.ToolBox(store, template_dir=str(tmp_path / "templates"))


@pytest.fixture
def fantomaas(store, tmp_path):
    return tools.ToolBox(store, mode="orchestrator", template_dir=str(tmp_path / "templates"))


def ok(r):
    assert r["status"] == "ok", r
    return r


def code(r):
    assert r["status"] == "error", r
    return r["error_code"]


def engine(store, sid="working"):
    return ce.compute_all(store.get_state(sid))


def row(table, label):
    return next(r for r in table["rows"] if r["label"] == label)


# ---- catalogue and dispatcher ----------------------------------------------------- #

def test_fourteen_tools_with_valid_schemas():
    assert len(tools.TOOL_SPECS) == 14
    for name, spec in tools.TOOL_SPECS.items():
        sch = spec["input_schema"]
        assert sch["type"] == "object" and set(sch["required"]) <= set(sch["properties"]), name
        assert spec["description"]
        assert hasattr(tools.ToolBox, name)


def test_visible_tools_follow_the_ship_gate_and_permissions(box, fantomaas):
    assert len(box.visible_tools()) == 13                       # optimizer hidden until step 8
    assert "find_optimal_mix" not in {t["name"] for t in box.visible_tools()}
    assert "manage_baseline_refresh" not in {t["name"] for t in fantomaas.visible_tools()}
    assert code(box.call("find_optimal_mix", {"objective": "group_cm"})) == "tool_unavailable"


@pytest.mark.parametrize("name, args, expected", [
    ("nope", {}, "invalid_path"),
    ("get_pnl", {}, "missing_required_arg"),
    ("get_pnl", {"scope": "galactic"}, "schema_validation_error"),
    ("get_pnl", {"scope": "standalone", "company": "EZDK", "colour": "red"}, "schema_validation_error"),
    ("get_pnl", {"scope": "standalone", "company": 7}, "schema_validation_error"),
    ("get_pnl", {"scope": "standalone"}, "missing_required_arg"),
    ("run_sensitivity", {"input_path": "dri.EZDK.mrmr", "variation": {"type": "range", "from": 1, "to": 2},
                         "output_metric": "ebt", "company": "EZDK"}, "schema_validation_error"),
    ("get_cost_sheet", {"stage": "finished_product", "company": "EZDK"}, "missing_required_arg"),
    ("read_state", {"state_id": "baseline_1999-01"}, "baseline_not_found"),
    ("read_state", {"state_id": "no-such-scenario"}, "scenario_not_found"),
])
def test_bad_calls_return_structured_errors(box, name, args, expected):
    r = box.call(name, args)
    assert code(r) == expected and r["error_message"]


@pytest.mark.parametrize("args", [
    {"stage": "dri", "company": "EFS"}, {"stage": "dri", "company": "ESR"},
    {"stage": "billet", "company": "ERM"},
    {"stage": "finished_product", "company": "ESR", "product": "HRC"},
    {"stage": "finished_product", "company": "EFS", "product": "Wire Rod"},
])
def test_structural_non_existence(box, args):
    r = box.call("get_cost_sheet", args)
    assert code(r) == "structural_non_existence" and "does not produce" in r["error_message"]
    assert code(box.call("get_sales_report", {"company": "ESR", "product": "HRC"})) == "structural_non_existence"


def test_all_companies_shows_producers_only_with_footnote(box):
    t = ok(box.call("get_cost_sheet", {"stage": "dri", "company": "all"}))["table"]
    assert t["columns"] == ["EZDK", "ERM"]
    assert any("EFS and ESR do not produce DRI" in f for f in t["footnotes"])
    t = ok(box.call("get_cost_sheet", {"stage": "finished_product", "company": "all", "product": "HRC"}))["table"]
    assert t["columns"] == ["EZDK", "EFS"]


# ---- every number is the engine's -------------------------------------------------- #

def test_cost_sheets_match_engine(box, store):
    o = engine(store)
    fx = store.get_state()["fx_egp_per_usd"]
    t = ok(box.call("get_cost_sheet", {"stage": "billet", "company": "all"}))["table"]
    assert row(t, "Total variable manufacturing cost")["values"]["EFS"] == o["billet"]["EFS"]["total_variable_cost"]
    t2 = ok(box.call("get_cost_sheet", {"stage": "billet", "company": "all", "currency": "EGP"}))["table"]
    assert row(t2, "Total variable manufacturing cost")["values"]["EFS"] == pytest.approx(
        fx * o["billet"]["EFS"]["total_variable_cost"])
    t = ok(box.call("get_cost_sheet", {"stage": "finished_product", "company": "all", "product": "Rebar"}))["table"]
    assert row(t, "Sc1 — in-house billet · Total variable manufacturing cost")["values"]["ERM"] == \
        o["rebar"]["ERM"]["sc1"]["total_variable_cost"]
    t = ok(box.call("get_cost_sheet", {"stage": "finished_product", "company": "all", "product": "Rebar",
                                       "view": "detailed"}))["table"]
    assert row(t, "Rebar variable cost")["values"] == {
        "EZDK": o["detail"]["EZDK Rebar"]["s3.variable_cost"], "EFS": o["detail"]["EFS Rebar"]["s3.variable_cost"],
        "ERM": o["detail"]["ERM Rolling"]["variable_cost_usd_t"], "ESR": o["detail"]["ESR Rebar"]["s3.variable_cost"]}


def test_dri_detailed_adds_up(box, store):
    o = engine(store)
    t = ok(box.call("get_cost_sheet", {"stage": "dri", "company": "all", "view": "detailed", "currency": "EGP"}))["table"]
    lines = [r for r in t["rows"] if r["label"] in ("IOP (MRMR × landed)", "Electricity", "Natural gas", "Oxygen",
                                                   "Nitrogen", "Water", "Chemicals", "Spare parts",
                                                   "External services", "Other")]
    for co in ("EZDK", "ERM"):
        assert sum(r["values"][f"{co} cost/t"] for r in lines) == pytest.approx(o["dri"][co]["variable_cost_le_t"])
        for r in lines[1:]:          # consumption × unit price = cost (natural gas priced per Nm³)
            v = r["values"]
            assert v[f"{co} consumption"] * v[f"{co} unit price"] == pytest.approx(v[f"{co} cost/t"])


def test_pnl_matches_engine(box, store):
    o = engine(store)
    t = ok(box.call("get_pnl", {"scope": "consolidated"}))["table"]
    assert t["columns"] == ["EZDK", "EFS", "ERM", "ESR", "Eliminations", "Consolidated"]
    con = o["consolidated"]["usd_monthly"]
    assert row(t, "EBT")["values"]["Consolidated"] == con["consolidated"]["ebt"]
    assert row(t, "Revenue")["values"]["Eliminations"] == con["eliminations"]["total_value"] < 0
    t = ok(box.call("get_pnl", {"scope": "standalone", "company": "ERM"}))["table"]
    assert t["columns"] == ["Rebar", "DRI (intercompany)", "Sub-Total"]
    assert row(t, "Contribution margin")["values"]["DRI (intercompany)"] == \
        o["pnl"]["usd_monthly"]["columns"]["ERM/DRI"]["contribution_margin"]
    le = ok(box.call("get_pnl", {"scope": "standalone", "company": "EZDK", "currency": "EGP", "period": "annual"}))["table"]
    assert row(le, "EBT")["values"]["Sub-Total"] == o["pnl"]["le_annual"]["columns"]["EZDK/Sub-Total"]["ebt"]
    assert row(le, "EBT")["unit"] == "M LE"


def test_sales_and_market_share(box, store):
    t = ok(box.call("get_sales_report", {"company": "all", "product": "Rebar"}))["table"]
    assert row(t, "Local sales")["values"]["EZDK Rebar"] == 100
    r = box.call("get_sales_report", {"company": "all", "product": "HRC", "view": "market_share"})
    assert r["status"] == "needs_input" and r["field"] == "hrc_total_local_market_kt"
    t = ok(box.call("get_sales_report", {"company": "all", "product": "HRC", "view": "market_share",
                                         "hrc_total_local_market_kt": 150}))["table"]
    assert row(t, "Ezz Steel market share")["values"]["HRC"] == pytest.approx(52.5 / 150 * 100)
    assert store.read("total_local_market_kt.HRC") is None          # never stored
    t = ok(box.call("get_sales_report", {"company": "EZDK", "product": "all", "view": "sales_mix"}))["table"]
    assert set(t["columns"]) == {"EZDK Rebar", "EZDK Wire Rod", "EZDK HRC"}


def test_fixed_cost_and_production_and_matrix(box, store):
    o = engine(store)
    t = ok(box.call("get_fixed_cost_report", {"company": "all"}))["table"]
    assert row(t, "Total with depreciation")["values"]["ERM"] == o["fixed_cost"]["usd_m"]["ERM"]["total_with_dep"]
    t = ok(box.call("get_fixed_cost_report", {"company": "ERM", "view": "per_unit"}))["table"]
    assert set(t["columns"]) == {"ERM/Rebar", "ERM/DRI"}
    r = ok(box.call("get_production_report", {"company": "all", "view": "full_plan"}))
    summary = next(x for x in r["tables"] if x["title"] == "Monthly production summary")
    assert row(summary, "DRI")["values"]["Total"] == o["production"]["summary_total"]["dri"]
    assert r["flow_diagram"]["mermaid"].startswith("flowchart LR") and r["flow_diagram"]["edges"]
    assert any(w["check"] == "capacity_within_limits" for w in r["warnings"])
    t = ok(box.call("get_tradeoff_matrix", {"buyer": "ERM"}))["table"]
    assert row(t, "Minimum")["values"]["ERM"] == o["billet"]["tradeoff"]["minimum"]["ERM"]
    assert row(t, "Cheapest source")["values"]["ERM"] == "EZDK"


# ---- analysis ------------------------------------------------------------------------- #

def test_sensitivity(box, store):
    base = engine(store)["pnl"]["usd_monthly"]["columns"]["ESR/Sub-Total"]["ebt"]
    r = ok(box.call("run_sensitivity", {"input_path": "sales.ESR.Rebar.local_price_le_t",
                                        "variation": {"type": "delta_pct", "values": [-10, 0, 10, 20]},
                                        "output_metric": "ebt", "company": "ESR"}))
    assert r["points"][1]["output"] == pytest.approx(base)
    assert r["zero_crossings"] and r["points"][0]["output"] < r["points"][-1]["output"]
    assert store.get_state() == store.get_state("baseline_2026-04")   # nothing stored
    r = ok(box.call("run_sensitivity", {"input_path": "billet.companies.EZDK.blend_pct.dri",
                                        "variation": {"type": "delta_abs", "values": [-5, 0]},
                                        "output_metric": "variable_cost_per_ton", "company": "EZDK",
                                        "product": "Rebar"}))
    assert [p["status"] for p in r["points"]] == ["integrity_failed", "ok"]   # flagged, not dropped
    r = ok(box.call("run_sensitivity", {"input_path": "sales.EZDK.Rebar.local_kt",
                                        "variation": {"type": "range", "from": -10, "to": 10, "steps": 3},
                                        "output_metric": "cm", "company": "EZDK"}))
    assert r["points"][0]["status"] == "invalid_input"
    assert code(box.call("run_sensitivity", {"input_path": "sourcing.ERM_rebar_billet",
                                             "variation": {"type": "delta_abs", "values": [1]},
                                             "output_metric": "ebt", "company": "ERM"})) == "schema_validation_error"


def test_update_and_compare_flag_a_sourcing_switch(box, fantomaas):
    r = ok(fantomaas.call("update_input", {"path": "dri.EZDK.iop_landed_usd_t", "value": 140}))
    assert r["sourcing_change"]["from"] == "EZDK" and r["sourcing_change"]["to"] == "EFS"
    assert "get_production_report" in r["affected_reports"]
    c = ok(box.call("compare_states", {"state_a_id": "baseline_2026-04", "state_b_id": "working",
                                       "output_focus": "cost_sheet"}))
    assert c["inputs"][0]["path"] == "dri.EZDK.iop_landed_usd_t"
    assert all("$/t" in o["figure"] for o in c["outputs"])
    assert c["sourcing_change"]["to"] == "EFS"
    c = ok(box.call("compare_states", {"state_a_id": "baseline_2026-04", "state_b_id": "working",
                                       "comparison_type": "inputs"}))
    assert "outputs" not in c


def test_update_batch_and_errors_via_tool(box):
    r = ok(box.call("update_input", {"changes": [
        {"path": "billet.companies.EZDK.blend_pct.dri", "value": 75},
        {"path": "billet.companies.EZDK.blend_pct.imported_scrap", "value": 21}]}))
    assert len(r["changes"]) == 2
    assert code(box.call("update_input", {"path": "billet.companies.EZDK.blend_pct.dri", "value": 70})) == \
        "integrity_check_failed"
    assert code(box.call("update_input", {})) == "missing_required_arg"


# ---- lifecycle --------------------------------------------------------------------- #

def test_commit_permissions(box, fantomaas):
    ok(fantomaas.call("commit_state", {"action": "save_as_scenario", "name": "f"}))
    assert code(fantomaas.call("commit_state", {"action": "promote_to_baseline", "month_label": "2026-05"})) == \
        "permission_denied"
    assert code(fantomaas.call("manage_baseline_refresh", {"action": "generate_template"})) == "permission_denied"
    assert ok(box.call("commit_state", {"action": "promote_to_baseline", "month_label": "2026-05"}))["baseline_id"] \
        == "baseline_2026-05"


def test_reports_refuse_a_state_failing_integrity(box, store, monkeypatch):
    import excel_io
    legacy = excel_io.extract_state(REFERENCE_WORKBOOK)       # blends ≠ 100, typed-in matrix cell
    monkeypatch.setattr(store, "get_state", lambda sid="working": legacy)
    assert code(box.call("get_pnl", {"scope": "consolidated"})) == "integrity_check_failed"


def test_template_has_appendix_d_layout_and_covers_every_input(box, store):
    r = ok(box.call("manage_baseline_refresh", {"action": "generate_template"}))
    wb = openpyxl.load_workbook(r["file_path"])
    assert wb.sheetnames == list(rt.SHEETS)
    seen = []
    for name in rt.SHEETS:
        ws = wb[name]
        header = next(rr for rr in range(1, ws.max_row + 1) if ws.cell(rr, 1).value == "Path")
        assert [ws.cell(header, c).value for c in range(1, 8)] == list(rt.HEADER)
        seen += [ws.cell(rr, 1).value for rr in range(header + 1, ws.max_row + 1)
                 if isinstance(ws.cell(rr, 1).value, str) and ws.cell(rr, 1).value]
    expected = [p for p in schema.leaves(store.get_state()) if rt.sheet_for(p)]
    assert sorted(seen) == sorted(expected) and len(seen) == len(set(seen))
    assert all(ws.cell(rr, rt.COL_THIS).value is None for ws in [wb["Sales"]]
               for rr in range(2, ws.max_row + 1) if ws.cell(rr, 1).value)   # empty by default


def _fill(path, values: dict, month=None):
    wb = openpyxl.load_workbook(path)
    for name in rt.SHEETS:
        ws = wb[name]
        for r in range(1, ws.max_row + 1):
            p = ws.cell(r, 1).value
            if p in values:
                ws.cell(r, rt.COL_THIS, values[p])
            if p == "Month label (YYYY-MM)" and month:
                ws.cell(r, 2, month)
    wb.save(path)


def test_refresh_round_trip_xlsx(box, store):
    path = ok(box.call("manage_baseline_refresh", {"action": "generate_template"}))["file_path"]
    _fill(path, {"sales.EZDK.Rebar.local_kt": 115, "fx_egp_per_usd": 16.5,
                 "detail.ESR Rebar.stage1.prices_usd.electrodes_kg": 230000}, month="2026-05")
    r = ok(box.call("manage_baseline_refresh", {"action": "commit", "source": "xlsx", "file_path": path}))
    assert r["baseline_id"] == "baseline_2026-05" and r["carried_forward_from"] == "baseline_2026-04"
    assert {c["path"] for c in r["entered_this_month"]} == {
        "sales.EZDK.Rebar.local_kt", "fx_egp_per_usd", "detail.ESR Rebar.stage1.prices_usd.electrodes_kg"}
    new = store.get_state("baseline_2026-05")
    assert new["detail"]["ESR Rebar"]["stage1"]["prices_usd"]["electrodes_kg"] == {"le": 230000.0}
    assert new["sales"]["EZDK"]["Rebar"]["local_kt"] == 115
    assert {d["path"] for d in r["diff_vs_previous"]["inputs"]} == {c["path"] for c in r["entered_this_month"]}
    # the same month again needs explicit overwrite
    assert code(box.call("manage_baseline_refresh", {"action": "commit", "source": "xlsx",
                                                     "file_path": path})) == "baseline_exists"


def test_refresh_rejects_bad_template_values(box, store):
    path = ok(box.call("manage_baseline_refresh", {"action": "generate_template"}))["file_path"]
    _fill(path, {"billet.companies.EFS.blend_pct.dri": 50}, month="2026-05")
    assert code(box.call("manage_baseline_refresh", {"action": "commit", "source": "xlsx",
                                                     "file_path": path})) == "integrity_check_failed"
    _fill(path, {"billet.companies.EFS.blend_pct.dri": "lots"}, month="2026-05")
    assert code(box.call("manage_baseline_refresh", {"action": "commit", "source": "xlsx",
                                                     "file_path": path})) == "schema_validation_error"
    assert [b["id"] for b in store.list_states("baselines")] == ["baseline_2026-04"]


def test_refresh_from_csv_and_payload(box, store, tmp_path):
    p = tmp_path / "m.csv"
    with open(p, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["path", "value", "comment"])
        w.writerow(["dri.ERM.iop_landed_usd_t", "150", "new contract"])
        w.writerow(["sourcing.ERM_rebar_billet", "market", ""])
    r = ok(box.call("manage_baseline_refresh", {"action": "commit", "source": "csv", "file_path": str(p),
                                                "month_label": "2026-05"}))
    assert store.read("sourcing.ERM_rebar_billet", "baseline_2026-05") == "market"
    with open(p, "w") as f:
        f.write("path,value\nsales.ESR.HRC.local_kt,5\n")
    assert code(box.call("manage_baseline_refresh", {"action": "commit", "source": "csv", "file_path": str(p),
                                                     "month_label": "2026-06"})) == "structural_non_existence"
    payload = store.get_state("baseline_2026-05")
    payload["fx_egp_per_usd"] = 17.0
    r = ok(box.call("manage_baseline_refresh", {"action": "commit", "source": "payload", "payload": payload,
                                                "month_label": "2026-06"}))
    assert r["previous_baseline_id"] == "baseline_2026-05"
    assert code(box.call("manage_baseline_refresh", {"action": "commit", "source": "xlsx",
                                                     "month_label": "2026-07"})) == "missing_required_arg"


def test_every_tool_returns_json_serialisable_results(box):
    calls = [("read_state", {"path": "sales"}), ("list_states", {}),
             ("get_cost_sheet", {"stage": "billet", "company": "all", "view": "detailed", "include_chart": True}),
             ("get_pnl", {"scope": "consolidated", "include_chart": True}),
             ("get_sales_report", {"company": "all", "product": "all", "include_chart": True}),
             ("get_fixed_cost_report", {"company": "all", "include_chart": True}),
             ("get_production_report", {"company": "all", "view": "full_plan"}),
             ("get_tradeoff_matrix", {}),
             ("run_sensitivity", {"input_path": "fx_egp_per_usd", "variation": {"type": "range", "from": 15,
                                  "to": 18, "steps": 4}, "output_metric": "ebt", "company": "EZDK",
                                  "include_chart": True}),
             ("compare_states", {"state_a_id": "baseline_2026-04", "state_b_id": "working"})]
    for name, args in calls:
        json.dumps(ok(box.call(name, args)))
